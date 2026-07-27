#!/usr/bin/env python3
"""Generate DOCX from HTML using python-docx with Japanese styling.

Usage:
    python generate_docx.py --input <html_file> --output <output_dir>
                            --skill-name <name>
                            [--template <docx_template>]

Parses the HTML document produced by the skill-to-document agent workflow
and renders it as a professionally formatted DOCX with book-quality styling.
"""

import argparse
import os
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


# ---------------------------------------------------------------------------
# Style configuration – book-quality Japanese document
# ---------------------------------------------------------------------------

_FONT_BODY = "游ゴシック"
_FONT_BODY_FALLBACK = "Hiragino Sans"
_FONT_HEADING = "游ゴシック"
_FONT_HEADING_FALLBACK = "Hiragino Sans"
_FONT_MONO = "Consolas"

_COLOR_HEADING = RGBColor(0x0F, 0x1B, 0x3D)
_COLOR_ACCENT = RGBColor(0x1B, 0x3A, 0x6B)
_COLOR_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
_COLOR_SECONDARY = RGBColor(0x4A, 0x4A, 0x5A)
_COLOR_TABLE_HEADER = RGBColor(0x1B, 0x3A, 0x6B)
_COLOR_TABLE_HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
_COLOR_TABLE_ALT = RGBColor(0xF5, 0xF7, 0xFA)
_COLOR_CODE_BG = "F0F2F6"


def _set_cell_shading(cell, color_hex: str):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {qn("w:fill"): color_hex, qn("w:val"): "clear"},
    )
    shading.append(shading_elem)


# ---------------------------------------------------------------------------
# HTML to DOCX parser
# ---------------------------------------------------------------------------

class _State:
    def __init__(self):
        self.bold = False
        self.italic = False
        self.in_code = False
        self.in_pre = False
        self.in_title_page = False
        self.in_toc = False
        self.in_table = False
        self.in_thead = False
        self.in_th = False
        self.in_td = False
        self.in_li = False
        self.in_blockquote = False
        self.heading_level = 0
        self.list_depth = 0
        self.list_type: list[str] = []
        self.list_counters: list[int] = []
        self.table_rows: list[list[str]] = []
        self.current_row: list[str] = []
        self.current_cell = ""
        self.link_href = ""
        self.pending_text = ""
        self.pre_text = ""


class HTMLToDocxParser(HTMLParser):
    """Parse skill-to-document HTML and render to python-docx Document."""

    def __init__(self, doc: Document, base_dir: str):
        super().__init__()
        self.doc = doc
        self.base_dir = base_dir
        self.s = _State()
        self._current_paragraph = None
        self._setup_styles()

    def _setup_styles(self):
        """Configure document styles for book-quality output."""
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = _FONT_BODY
        font.size = Pt(10.5)
        font.color.rgb = _COLOR_TEXT
        pf = style.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.6

        # Set East Asian font
        rpr = style.element.get_or_add_rPr()
        ea = rpr.makeelement(qn("w:rFonts"), {qn("w:eastAsia"): _FONT_BODY})
        rpr.append(ea)

        for level in range(1, 7):
            style_name = f"Heading {level}"
            if style_name in self.doc.styles:
                hs = self.doc.styles[style_name]
                hs.font.name = _FONT_HEADING
                hs.font.color.rgb = _COLOR_HEADING
                hs.font.bold = True
                sizes = {1: 20, 2: 16, 3: 14, 4: 12, 5: 11, 6: 10}
                hs.font.size = Pt(sizes.get(level, 10))
                hs.paragraph_format.space_before = Pt(18 if level <= 2 else 12)
                hs.paragraph_format.space_after = Pt(6)
                hs_rpr = hs.element.get_or_add_rPr()
                hs_ea = hs_rpr.makeelement(
                    qn("w:rFonts"), {qn("w:eastAsia"): _FONT_HEADING}
                )
                hs_rpr.append(hs_ea)

    def _add_paragraph(self, text="", style=None):
        self._current_paragraph = self.doc.add_paragraph(text, style=style)
        return self._current_paragraph

    def _add_run(
        self, text, bold=False, italic=False,
        font_name=None, font_size=None, color=None,
    ):
        if self._current_paragraph is None:
            self._add_paragraph()
        run = self._current_paragraph.add_run(text)
        run.bold = bold or self.s.bold
        run.italic = italic or self.s.italic
        if font_name:
            run.font.name = font_name
        if font_size:
            run.font.size = font_size
        if color:
            run.font.color.rgb = color
        return run

    def _flush_text(self):
        text = self.s.pending_text
        if not text:
            return
        self.s.pending_text = ""

        if self.s.in_table:
            self.s.current_cell += text
            return

        if self.s.in_pre:
            self.s.pre_text += text
            return

        if self.s.heading_level:
            self._add_paragraph(text, style=f"Heading {self.s.heading_level}")
            return

        if self.s.in_title_page:
            p = self._add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(24)
            run.font.color.rgb = _COLOR_HEADING
            run.font.name = _FONT_HEADING
            return

        if self.s.in_code and not self.s.in_pre:
            self._add_run(text, font_name=_FONT_MONO, font_size=Pt(9))
            return

        self._add_run(text)

    # -- tag handlers --

    def handle_starttag(self, tag, attrs):
        attr = dict(attrs)
        cls = attr.get("class", "")

        if tag == "div" and "title-page" in cls:
            self.s.in_title_page = True
            return

        if tag == "div" and "toc" in cls:
            self.s.in_toc = True
            return

        if tag == "div" and "chapter" in cls:
            self.doc.add_page_break()
            return

        if tag == "div" and "appendix" in cls:
            self.doc.add_page_break()
            return

        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush_text()
            self.s.heading_level = int(tag[1])
            self._current_paragraph = None
            return

        if tag == "p":
            self._flush_text()
            self._current_paragraph = None
            return

        if tag in ("strong", "b"):
            self._flush_text()
            self.s.bold = True
            return

        if tag in ("em", "i"):
            self._flush_text()
            self.s.italic = True
            return

        if tag == "code":
            self._flush_text()
            self.s.in_code = True
            return

        if tag == "pre":
            self._flush_text()
            self.s.in_pre = True
            self.s.pre_text = ""
            return

        if tag in ("ul", "ol"):
            self.s.list_depth += 1
            self.s.list_type.append(tag)
            self.s.list_counters.append(0)
            return

        if tag == "li":
            self._flush_text()
            self.s.in_li = True
            self.s.list_counters[-1] += 1
            if self.s.list_type[-1] == "ol":
                prefix = f"{self.s.list_counters[-1]}. "
            else:
                prefix = "• "
            style = "List Bullet" if self.s.list_type[-1] == "ul" else "List Number"
            p = self._add_paragraph(style=style)
            p.paragraph_format.left_indent = Cm(1.0 * self.s.list_depth)
            self._current_paragraph = p
            # Add prefix as run for ordered lists
            if self.s.list_type[-1] == "ol":
                p.clear()
                p.add_run(prefix)
            return

        if tag == "blockquote":
            self._flush_text()
            self.s.in_blockquote = True
            return

        if tag == "table":
            self._flush_text()
            self.s.in_table = True
            self.s.table_rows = []
            return

        if tag == "thead":
            self.s.in_thead = True
            return

        if tag == "tr":
            self.s.current_row = []
            return

        if tag in ("th", "td"):
            self.s.in_th = tag == "th"
            self.s.in_td = tag == "td"
            self.s.current_cell = ""
            return

        if tag == "img":
            self._flush_text()
            src = attr.get("src", "")
            img_path = os.path.join(self.base_dir, src) if src else ""
            if img_path and os.path.exists(img_path):
                try:
                    self.doc.add_picture(img_path, width=Cm(15))
                except Exception:
                    self._add_paragraph(f"[画像: {src} — 読み込みエラー]")
            else:
                self._add_paragraph(f"[画像: {src} — 取得不可]")
            return

        if tag == "hr":
            self._flush_text()
            p = self._add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            # Add a border line
            pPr = p._element.get_or_add_pPr()
            pBdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pBdr.makeelement(
                qn("w:bottom"),
                {
                    qn("w:val"): "single",
                    qn("w:sz"): "4",
                    qn("w:space"): "1",
                    qn("w:color"): "C5CDD9",
                },
            )
            pBdr.append(bottom)
            pPr.append(pBdr)
            return

        if tag == "br":
            if self._current_paragraph:
                self._current_paragraph.add_run("\n")
            return

    def handle_endtag(self, tag):
        if tag == "div":
            if self.s.in_title_page:
                self.s.in_title_page = False
                self.doc.add_page_break()
            if self.s.in_toc:
                self.s.in_toc = False
            return

        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush_text()
            self.s.heading_level = 0
            return

        if tag == "p":
            self._flush_text()
            self._current_paragraph = None
            return

        if tag in ("strong", "b"):
            self._flush_text()
            self.s.bold = False
            return

        if tag in ("em", "i"):
            self._flush_text()
            self.s.italic = False
            return

        if tag == "code":
            self._flush_text()
            self.s.in_code = False
            return

        if tag == "pre":
            # Render collected pre text as a code block
            code_text = self.s.pre_text.strip()
            self.s.in_pre = False
            self.s.in_code = False
            self.s.pre_text = ""
            if code_text:
                p = self._add_paragraph()
                run = p.add_run(code_text)
                run.font.name = _FONT_MONO
                run.font.size = Pt(9)
                # Background shading for code block
                pPr = p._element.get_or_add_pPr()
                shd = pPr.makeelement(
                    qn("w:shd"),
                    {qn("w:fill"): _COLOR_CODE_BG, qn("w:val"): "clear"},
                )
                pPr.append(shd)
            return

        if tag in ("ul", "ol"):
            if self.s.list_depth > 0:
                self.s.list_depth -= 1
                self.s.list_type.pop() if self.s.list_type else None
                self.s.list_counters.pop() if self.s.list_counters else None
            return

        if tag == "li":
            self._flush_text()
            self.s.in_li = False
            self._current_paragraph = None
            return

        if tag == "blockquote":
            self._flush_text()
            self.s.in_blockquote = False
            return

        if tag == "thead":
            self.s.in_thead = False
            return

        if tag in ("th", "td"):
            self.s.current_row.append(self.s.current_cell.strip())
            self.s.current_cell = ""
            self.s.in_th = False
            self.s.in_td = False
            return

        if tag == "tr":
            if self.s.current_row:
                self.s.table_rows.append(self.s.current_row)
            return

        if tag == "table":
            self._render_table()
            self.s.in_table = False
            self.s.table_rows = []
            return

    def handle_data(self, data):
        if self.s.in_table:
            self.s.current_cell += data
        elif self.s.in_pre:
            self.s.pre_text += data
        else:
            self.s.pending_text += data

    def _render_table(self):
        rows = self.s.table_rows
        if not rows:
            return
        num_cols = max(len(r) for r in rows)
        if num_cols == 0:
            return

        table = self.doc.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for i, row_data in enumerate(rows):
            row = table.rows[i]
            is_header = i == 0
            for j in range(num_cols):
                cell = row.cells[j]
                cell.text = row_data[j] if j < len(row_data) else ""
                p = cell.paragraphs[0]
                if is_header:
                    run = p.runs[0] if p.runs else p.add_run(cell.text)
                    run.bold = True
                    run.font.color.rgb = _COLOR_TABLE_HEADER_TEXT
                    run.font.size = Pt(9)
                    _set_cell_shading(cell, "1B3A6B")
                else:
                    for run in p.runs:
                        run.font.size = Pt(9)
                    if i % 2 == 0:
                        _set_cell_shading(cell, "F5F7FA")

        self._current_paragraph = None


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Generate DOCX from HTML")
    parser.add_argument("--input", required=True, help="Input HTML file path")
    parser.add_argument("--output", required=True, help="Output directory")
    parser.add_argument(
        "--skill-name", required=True,
        help="Skill name for output filename",
    )
    parser.add_argument("--template", default=None, help="Optional DOCX template file")
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    output_dir = Path(args.output).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    if not input_path.exists():
        print(f"Error: Input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    html_content = input_path.read_text(encoding="utf-8")
    base_dir = str(input_path.parent)

    if args.template and Path(args.template).exists():
        doc = Document(args.template)
    else:
        doc = Document()

    # Page setup – A4
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)

    parser_inst = HTMLToDocxParser(doc, base_dir)
    parser_inst.feed(html_content)
    parser_inst._flush_text()

    output_path = output_dir / f"{args.skill_name}.docx"
    doc.save(str(output_path))
    print(f"DOCX generated: {output_path}")


if __name__ == "__main__":
    main()
